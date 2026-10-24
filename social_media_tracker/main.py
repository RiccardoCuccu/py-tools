#!/usr/bin/env python3
"""
Social Media Tracker - Track Instagram and YouTube follower counts on a schedule and export to Excel.

Runs continuously, checking follower counts at configured intervals (daily, last day of
month, or a specific weekday) and appending results to a shared Excel file.

Usage:
    python main.py

Notes:
    - Profiles are configured in social_profiles.json (not committed).
"""

import calendar
import json
import os
import re
import time
import traceback
from datetime import datetime
from typing import Dict, Optional, Any

import openpyxl
import schedule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import scraper_instagram, scraper_threads, scraper_youtube

# General configuration - edit these values

# Debug Settings
DEBUG_MODE = False  # Set to False to disable detailed debug logs

# Configuration file name (not to be committed to git)
CONFIG_FILE = "social_profiles.json"
EXCEL_FILE = "social_followers_history.xlsx"

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(SCRIPT_DIR, CONFIG_FILE)
EXCEL_FILE = os.path.join(SCRIPT_DIR, EXCEL_FILE)
SCHEDULER_POLL_INTERVAL_SECONDS = 60

# Schedule validation
VALID_SCHEDULE_MODES = ("daily", "weekly", "monthly")
TIME_PATTERN = re.compile(r"([01]\d|2[0-3]):[0-5]\d")
DATE_FORMAT = "%Y-%m-%d"
ERROR_MARKER = "Error"
WEEKDAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

# Debug utilities

def debug_log(message: str) -> None:
    """Print debug message if DEBUG_MODE is enabled"""
    if DEBUG_MODE:
        timestamp = datetime.now().strftime("%H:%M:%S")
        print(f"[DEBUG {timestamp}] {message}")

# Platform scrapers registry

# Registry of available scrapers - add new platforms here
SCRAPERS = {
    "instagram": scraper_instagram.get_followers,
    "threads": scraper_threads.get_followers,
    "youtube": scraper_youtube.get_followers,
}

# First run setup

def create_default_config() -> None:
    """Create default configuration file if it doesn't exist."""
    default_config = {
        "schedules": {
            "_comment_mode": "Options: 'daily', 'weekly', or 'monthly'",
            "mode": "monthly",
            "_comment_weekly_day": "0=Monday, 1=Tuesday, 2=Wednesday, 3=Thursday, 4=Friday, 5=Saturday, 6=Sunday (only used if mode is 'weekly')",
            "weekly_day": 0,
            "_comment_time": "Execution time in HH:MM format (24-hour)",
            "time": "23:00"
        },
        "profiles": {
            "instagram": {
                "url": "",
                "enabled": False
            },
            "threads": {
                "url": "",
                "enabled": False
            },
            "youtube": {
                "url": "",
                "enabled": False
            }
        },
        "settings": {
            "retry_attempts": 3,
            "retry_delay_seconds": 5,
            "delay_between_checks_seconds": 2
        }
    }
    
    with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
        json.dump(default_config, f, indent=4)
    print(f"Created default config file: {CONFIG_FILE}")
    print("Please edit the file to add your profile URLs and enable tracking.")

def load_config() -> Dict[str, Any]:
    """Load configuration from JSON file."""
    if not os.path.exists(CONFIG_FILE):
        create_default_config()
        raise FileNotFoundError(f"Please configure {CONFIG_FILE} first")
    
    try:
        with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
            config = json.load(f)
    except json.JSONDecodeError as e:
        raise ValueError(f"{CONFIG_FILE} is not valid JSON: {e}") from e

    validate_config(config)
    return config

def validate_config(config: Dict[str, Any]) -> None:
    """Raise ValueError with a clear message if the schedule/profile settings are unusable."""
    sched_config = config.get("schedules")
    if not isinstance(sched_config, dict) or not isinstance(config.get("profiles"), dict):
        raise ValueError("Config must contain 'schedules' and 'profiles' sections")

    mode = sched_config.get("mode")
    if mode not in VALID_SCHEDULE_MODES:
        raise ValueError(f"schedules.mode must be one of {VALID_SCHEDULE_MODES}, got {mode!r}")

    time_value = sched_config.get("time")
    if not isinstance(time_value, str) or not TIME_PATTERN.fullmatch(time_value):
        raise ValueError(f"schedules.time must be HH:MM (24-hour), got {time_value!r}")

    if mode == "weekly":
        weekly_day = sched_config.get("weekly_day")
        if not isinstance(weekly_day, int) or isinstance(weekly_day, bool) \
                or not 0 <= weekly_day < len(WEEKDAY_NAMES):
            raise ValueError(f"schedules.weekly_day must be 0-6, got {weekly_day!r}")

# Data collection

def collect_stats(config: Dict[str, Any]) -> Dict[str, Dict[str, Any]]:
    """Collect follower stats from all enabled platforms."""
    stats = {}
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    settings = config.get("settings", {})
    max_retries = settings.get("retry_attempts", 3)
    retry_delay = settings.get("retry_delay_seconds", 5)
    check_delay = settings.get("delay_between_checks_seconds", 2)
    
    for platform, profile_info in config["profiles"].items():
        if not profile_info.get("enabled", False) or not profile_info.get("url"):
            continue
        
        if platform not in SCRAPERS:
            print(f"Warning: No scraper found for platform '{platform}', skipping...")
            continue
        
        url = profile_info["url"]
        print(f"Checking {platform}: {url}")

        try:
            scraper_func = SCRAPERS[platform]

            result = scraper_func(url, max_retries=max_retries, retry_delay=retry_delay, debug_mode=DEBUG_MODE)

            if result.get("followers") is not None:
                username_display = result.get("username", "Unknown")
                followers = result.get("followers")
                print(f"  OK {platform.title()}: {username_display} - {followers:,} followers")
            else:
                print(f"  FAILED {platform.title()}: Could not retrieve follower count")
            
            stats[platform] = {
                "username": result.get("username", ""),
                "url": url,
                "followers": result.get("followers"),
                "checked_at": timestamp
            }
            
            # Delay between checks to avoid rate limiting
            if check_delay > 0 and platform != list(config["profiles"].keys())[-1]:
                time.sleep(check_delay)
                
        except Exception as e:
            print(f"Error checking {platform}: {e}")
            stats[platform] = {
                "username": "",
                "url": url,
                "followers": None,
                "checked_at": timestamp
            }
    
    return stats

# Excel reporting

def last_numeric_followers(ws: Worksheet) -> Optional[int]:
    """Return the most recent numeric follower count in column D, skipping 'Error' rows."""
    for row in range(ws.max_row, 1, -1):
        value = ws.cell(row=row, column=4).value
        if isinstance(value, int) and not isinstance(value, bool):
            return value
    return None

def period_key(day: datetime, mode: str) -> tuple:
    """Return the schedule period a date belongs to: day, ISO week or month."""
    if mode == "monthly":
        return (day.year, day.month)
    if mode == "weekly":
        iso = day.isocalendar()
        return (iso[0], iso[1])
    return (day.year, day.month, day.day)

def current_period_row(ws: Worksheet, mode: str) -> Optional[int]:
    """Return the last row if it was recorded in the current period, else None.

    Rows are chronological, so only the last row can belong to the current period.
    Unparseable dates (hand-edited sheets) are treated as a different period.
    """
    if ws.max_row < 2:
        return None
    raw = ws.cell(row=ws.max_row, column=1).value
    try:
        recorded = datetime.strptime(str(raw), DATE_FORMAT)
    except ValueError:
        return None
    if period_key(recorded, mode) == period_key(datetime.now(), mode):
        return ws.max_row
    return None

def recorded_platforms(mode: str) -> set:
    """Return the sheet names that already hold a successful row for the current period."""
    if not os.path.exists(EXCEL_FILE):
        return set()
    wb = openpyxl.load_workbook(EXCEL_FILE, read_only=False)
    done = set()
    for ws in wb.worksheets:
        row = current_period_row(ws, mode)
        if row is not None and ws.cell(row=row, column=4).value != ERROR_MARKER:
            done.add(ws.title)
    return done

def append_to_excel_history(stats: Dict[str, Dict[str, Any]], mode: Optional[str] = None) -> None:
    """Append follower statistics to separate sheets per platform in Excel file.

    With a schedule mode, at most one row per platform is kept per period (day, ISO
    week or month): a period already holding a successful row is left untouched, and
    an 'Error' row of the current period is replaced by the new result.
    """
    date_str = datetime.now().strftime(DATE_FORMAT)
    time_str = datetime.now().strftime("%H:%M:%S")
    
    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
    else:
        wb = openpyxl.Workbook()
        # Remove default sheet
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
    
    for platform, data in stats.items():
        platform_name = platform.title()
        
        # Get or create sheet for this platform
        is_new_sheet = platform_name not in wb.sheetnames
        if platform_name in wb.sheetnames:
            ws = wb[platform_name]
        else:
            ws = wb.create_sheet(platform_name)

        if not isinstance(ws, Worksheet):
            continue

        if is_new_sheet:
            headers = ["Date", "Time", "Username", "Followers", "Change"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        last_row = ws.max_row

        # Get previous value for change calculation
        previous_followers = last_numeric_followers(ws)

        row = last_row + 1
        existing = current_period_row(ws, mode) if mode else None
        if existing is not None:
            if ws.cell(row=existing, column=4).value != ERROR_MARKER:
                continue
            # Replace the failed attempt of this period instead of adding a second row
            row = existing
            for col in range(1, 6):
                ws.cell(row=row, column=col).value = None
                ws.cell(row=row, column=col).font = Font()
        ws.cell(row=row, column=1, value=date_str)
        ws.cell(row=row, column=2, value=time_str)
        ws.cell(row=row, column=3, value=data["username"])
        
        if data["followers"] is not None:
            ws.cell(row=row, column=4, value=data["followers"])
            
            # Calculate change from previous check
            if previous_followers is not None:
                change = data["followers"] - previous_followers
                change_str = f"+{change}" if change > 0 else str(change)
                change_cell = ws.cell(row=row, column=5, value=change_str)
                
                if change > 0:
                    change_cell.font = Font(color="00AA00")  # Green
                elif change < 0:
                    change_cell.font = Font(color="FF0000")  # Red
        else:
            ws.cell(row=row, column=4, value=ERROR_MARKER)
        
        for col in range(1, 6):
            ws.column_dimensions[get_column_letter(col)].width = 15
    
    # Save to a temp file first so a crash mid-write cannot corrupt the history
    tmp_file = EXCEL_FILE + ".tmp"
    wb.save(tmp_file)
    os.replace(tmp_file, EXCEL_FILE)
    print(f"History updated: {EXCEL_FILE}")

# Scheduling

def is_last_day_of_month() -> bool:
    """Check if today is the last day of the current month."""
    today = datetime.now()
    _, last_day = calendar.monthrange(today.year, today.month)
    return today.day == last_day

def is_target_weekday(weekday: int) -> bool:
    """Check if today matches the target weekday."""
    return datetime.now().weekday() == weekday

def should_run_job(config: Dict[str, Any]) -> bool:
    """Check if job should run based on schedule configuration."""
    sched_config = config["schedules"]
    
    if sched_config["mode"] == "daily":
        return True
    elif sched_config["mode"] == "monthly":
        return is_last_day_of_month()
    elif sched_config["mode"] == "weekly":
        return is_target_weekday(sched_config["weekly_day"])
    
    return False

def job_runner() -> None:
    """Main job: collect stats and generate report."""
    config = load_config()
    
    if not should_run_job(config):
        return
    
    print(f"\n{'='*60}")
    print(f"Running follower check at {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"{'='*60}\n")
    
    # Startup and scheduled runs can fall in the same period: skip platforms already recorded
    mode = config["schedules"]["mode"]
    done = {name.lower() for name in recorded_platforms(mode)}
    pending = {k: v for k, v in config["profiles"].items() if k.lower() not in done}
    if done:
        print(f"Already recorded this period, skipping: {', '.join(sorted(done))}")
    stats = collect_stats({**config, "profiles": pending})

    print(f"\n{'='*60}")
    print("SUMMARY")
    print(f"{'='*60}")
    for platform, data in stats.items():
        if data.get("followers") is not None:
            print(f"{platform.title():12} : {data['username']:25} - {data['followers']:>8,} followers")
        else:
            print(f"{platform.title():12} : {data.get('username', 'N/A'):25} - {'ERROR':>8}")
    print(f"{'='*60}\n")
    
    append_to_excel_history(stats, mode)
    
    print(f"{'='*60}")
    print("Job completed successfully!")
    print(f"{'='*60}\n")

def safe_job_runner() -> None:
    """Run job_runner, reporting any failure instead of killing the scheduler loop."""
    try:
        job_runner()
    except Exception as e:
        # A locked Excel file, bad config edit or scraper bug must not stop the daemon
        print(f"\nJob failed: {type(e).__name__}: {e}")
        traceback.print_exc()

def setup_schedule(config: Dict[str, Any]) -> None:
    """Setup scheduled jobs based on configuration."""
    sched_config = config["schedules"]
    
    # Schedule to check every day at specified time
    schedule.every().day.at(sched_config["time"]).do(safe_job_runner)
    
    if sched_config["mode"] == "daily":
        print(f"Scheduled: Daily at {sched_config['time']}")
    elif sched_config["mode"] == "monthly":
        print(f"Scheduled: Monthly (last day) at {sched_config['time']}")
    elif sched_config["mode"] == "weekly":
        weekday_name = WEEKDAY_NAMES[sched_config["weekly_day"]]
        print(f"Scheduled: Weekly ({weekday_name}) at {sched_config['time']}")

# Main

def main() -> None:
    """Main function: load config, setup schedule, run forever."""
    try:
        config = load_config()
        print("\n" + "="*60)
        print("Social Media Follower Tracker Started!")
        print("="*60)
        print(f"Config loaded from: {CONFIG_FILE}")
        print(f"History file: {EXCEL_FILE}")
        print(f"Debug mode: {'ENABLED' if DEBUG_MODE else 'DISABLED'}")
        print(f"Available platforms: {', '.join(SCRAPERS.keys())}")
        
        setup_schedule(config)
        
        # Run immediately once on startup (if conditions met)
        print("\nRunning initial check...")
        safe_job_runner()

        print("\nWaiting for scheduled runs...")
        print("Press Ctrl+C to stop\n")
        
        while True:
            schedule.run_pending()
            time.sleep(SCHEDULER_POLL_INTERVAL_SECONDS)
            
    except KeyboardInterrupt:
        print("\n\nTracker stopped by user.")
    except FileNotFoundError as e:
        print(f"\n{e}")
    except Exception as e:
        print(f"\nError: {e}")
        traceback.print_exc()

if __name__ == "__main__":
    main()