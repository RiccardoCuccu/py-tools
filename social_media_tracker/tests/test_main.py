"""Tests for main: config creation/validation, scheduling rules, stat collection, Excel
history (change tracking, error rows, atomic save) and the failure-tolerant job runner.
Every file path is redirected to tmp_path and all scrapers are faked.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

import main as tracker


def make_config(**overrides) -> dict:
    config = {
        "schedules": {"mode": "daily", "weekly_day": 0, "time": "23:00"},
        "profiles": {"instagram": {"url": "https://i/x", "enabled": True}},
        "settings": {"retry_attempts": 1, "retry_delay_seconds": 0,
                     "delay_between_checks_seconds": 0},
    }
    config.update(overrides)
    return config


@pytest.fixture(autouse=True)
def redirect_paths(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Point every production path at tmp_path."""
    monkeypatch.setattr(tracker, "CONFIG_FILE", str(tmp_path / "social_profiles.json"))
    monkeypatch.setattr(tracker, "EXCEL_FILE", str(tmp_path / "history.xlsx"))
    monkeypatch.setattr(tracker.time, "sleep", lambda s: None)


class FixedDate(datetime):
    """datetime whose now() is controlled by the test."""

    fixed = datetime(2024, 2, 29, 12, 0, 0)

    @classmethod
    def now(cls, tz=None):
        return cls.fixed


# Config

def test_load_config_creates_default_and_raises_when_missing(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        tracker.load_config()
    created = json.loads((tmp_path / "social_profiles.json").read_text(encoding="utf-8"))
    assert created["profiles"]["instagram"]["enabled"] is False
    # The generated default is itself valid
    tracker.validate_config(created)


def test_load_config_reads_valid_file(tmp_path: Path) -> None:
    (tmp_path / "social_profiles.json").write_text(json.dumps(make_config()), encoding="utf-8")
    assert tracker.load_config()["schedules"]["mode"] == "daily"


def test_load_config_invalid_json_gives_clear_error(tmp_path: Path) -> None:
    (tmp_path / "social_profiles.json").write_text("{not json", encoding="utf-8")
    with pytest.raises(ValueError, match="not valid JSON"):
        tracker.load_config()


@pytest.mark.parametrize(
    "schedules",
    [
        {"mode": "hourly", "time": "10:00"},
        {"mode": "daily", "time": "25:00"},
        {"mode": "daily", "time": "9:00"},
        {"mode": "daily", "time": 900},
        {"mode": "daily"},
        {"mode": "weekly", "time": "10:00", "weekly_day": 7},
        {"mode": "weekly", "time": "10:00", "weekly_day": -1},
        {"mode": "weekly", "time": "10:00", "weekly_day": "1"},
        {"mode": "weekly", "time": "10:00", "weekly_day": True},
        {"mode": "weekly", "time": "10:00"},
    ],
)
def test_validate_config_rejects_bad_schedules(schedules: dict) -> None:
    with pytest.raises(ValueError):
        tracker.validate_config(make_config(schedules=schedules))


def test_validate_config_rejects_missing_sections() -> None:
    with pytest.raises(ValueError):
        tracker.validate_config({"schedules": {"mode": "daily", "time": "10:00"}})
    with pytest.raises(ValueError):
        tracker.validate_config({"profiles": {}})


def test_validate_config_accepts_all_modes() -> None:
    for mode in ("daily", "monthly"):
        tracker.validate_config(make_config(schedules={"mode": mode, "time": "00:00"}))
    tracker.validate_config(
        make_config(schedules={"mode": "weekly", "weekly_day": 6, "time": "23:59"})
    )


# Scheduling

def test_is_last_day_of_month(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(tracker, "datetime", FixedDate)
    FixedDate.fixed = datetime(2024, 2, 29)  # leap year
    assert tracker.is_last_day_of_month() is True
    FixedDate.fixed = datetime(2024, 2, 28)
    assert tracker.is_last_day_of_month() is False
    FixedDate.fixed = datetime(2023, 2, 28)
    assert tracker.is_last_day_of_month() is True


def test_should_run_job_by_mode(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(tracker, "datetime", FixedDate)
    FixedDate.fixed = datetime(2024, 3, 4)  # a Monday, not month end

    def cfg(mode: str, day: int = 0) -> dict:
        return {"schedules": {"mode": mode, "weekly_day": day, "time": "10:00"}}

    assert tracker.should_run_job(cfg("daily")) is True
    assert tracker.should_run_job(cfg("monthly")) is False
    assert tracker.should_run_job(cfg("weekly", 0)) is True
    assert tracker.should_run_job(cfg("weekly", 1)) is False
    assert tracker.should_run_job(cfg("bogus")) is False


# Collection

def test_collect_stats_success_and_missing_count(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[dict] = []

    def fake_ig(url, **kwargs):
        calls.append(kwargs)
        return {"username": "alice", "followers": 1234}

    monkeypatch.setattr(
        tracker, "SCRAPERS",
        {"instagram": fake_ig, "youtube": lambda url, **k: {"username": "", "followers": None}},
    )
    config = make_config(profiles={
        "instagram": {"url": "https://i/alice", "enabled": True},
        "youtube": {"url": "https://y/c", "enabled": True},
    })
    stats = tracker.collect_stats(config)

    assert stats["instagram"]["followers"] == 1234
    assert stats["instagram"]["username"] == "alice"
    assert stats["youtube"]["followers"] is None
    assert calls[0]["max_retries"] == 1


def test_collect_stats_skips_disabled_empty_and_unknown(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(tracker, "SCRAPERS", {"instagram": lambda url, **k: {"followers": 1}})
    config = make_config(profiles={
        "instagram": {"url": "", "enabled": True},
        "threads": {"url": "https://t/x", "enabled": True},
        "youtube": {"url": "https://y/x", "enabled": False},
    })
    assert tracker.collect_stats(config) == {}


def test_collect_stats_scraper_exception_becomes_failed_entry(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def boom(url, **kwargs):
        raise RuntimeError("blocked")

    monkeypatch.setattr(tracker, "SCRAPERS", {"instagram": boom})
    stats = tracker.collect_stats(make_config())
    assert stats["instagram"]["followers"] is None
    assert stats["instagram"]["url"] == "https://i/x"


def test_collect_stats_survives_unencodable_console(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A successful scrape must stay a success even with a non-ASCII username."""
    monkeypatch.setattr(
        tracker, "SCRAPERS", {"instagram": lambda url, **k: {"username": "ünï", "followers": 5}}
    )
    stats = tracker.collect_stats(make_config())
    assert stats["instagram"]["followers"] == 5


# Excel

def stat(followers, username="alice") -> dict:
    return {"instagram": {"username": username, "url": "u", "followers": followers,
                          "checked_at": "t"}}


def read_rows(path: str, sheet: str = "Instagram") -> list[tuple]:
    wb = openpyxl.load_workbook(path)
    return [tuple(c.value for c in row) for row in wb[sheet].iter_rows()]


def test_excel_creates_workbook_with_header_and_first_row(tmp_path: Path) -> None:
    tracker.append_to_excel_history(stat(100))
    rows = read_rows(tracker.EXCEL_FILE)

    assert rows[0] == ("Date", "Time", "Username", "Followers", "Change")
    assert rows[1][2:] == ("alice", 100, None)
    wb = openpyxl.load_workbook(tracker.EXCEL_FILE)
    assert wb.sheetnames == ["Instagram"]
    assert not list(tmp_path.glob("*.tmp"))


def test_excel_records_signed_change(tmp_path: Path) -> None:
    tracker.append_to_excel_history(stat(100))
    tracker.append_to_excel_history(stat(150))
    tracker.append_to_excel_history(stat(140))
    rows = read_rows(tracker.EXCEL_FILE)
    assert [r[3] for r in rows[1:]] == [100, 150, 140]
    assert [r[4] for r in rows[1:]] == [None, "+50", "-10"]


def test_excel_zero_change_and_error_rows() -> None:
    tracker.append_to_excel_history(stat(100))
    tracker.append_to_excel_history(stat(None))
    tracker.append_to_excel_history(stat(100))
    rows = read_rows(tracker.EXCEL_FILE)
    assert rows[2][3] == "Error"
    # The change is computed against the last real count, not the Error row
    assert rows[3][3:] == (100, "0")


def test_excel_first_value_after_error_has_no_change() -> None:
    tracker.append_to_excel_history(stat(None))
    tracker.append_to_excel_history(stat(10))
    rows = read_rows(tracker.EXCEL_FILE)
    assert rows[2][3:] == (10, None)


def test_excel_separate_sheet_per_platform() -> None:
    stats = stat(1)
    stats["threads"] = {"username": "t", "url": "u", "followers": 2, "checked_at": "t"}
    tracker.append_to_excel_history(stats)
    wb = openpyxl.load_workbook(tracker.EXCEL_FILE)
    assert wb.sheetnames == ["Instagram", "Threads"]


def test_excel_failed_save_keeps_previous_history(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    tracker.append_to_excel_history(stat(100))
    before = Path(tracker.EXCEL_FILE).read_bytes()

    def failing_replace(src, dst):
        raise PermissionError("locked")

    monkeypatch.setattr(tracker.os, "replace", failing_replace)
    with pytest.raises(PermissionError):
        tracker.append_to_excel_history(stat(200))
    assert Path(tracker.EXCEL_FILE).read_bytes() == before


# Job runner

def test_safe_job_runner_swallows_errors(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    def boom() -> None:
        raise PermissionError("excel is open")

    monkeypatch.setattr(tracker, "job_runner", boom)
    tracker.safe_job_runner()  # must not raise
    assert "Job failed" in capsys.readouterr().out


def test_job_runner_collects_and_writes_history(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    (tmp_path / "social_profiles.json").write_text(json.dumps(make_config()), encoding="utf-8")
    monkeypatch.setattr(
        tracker, "SCRAPERS", {"instagram": lambda url, **k: {"username": "a", "followers": 7}}
    )
    tracker.job_runner()
    assert read_rows(tracker.EXCEL_FILE)[1][3] == 7


def test_job_runner_does_nothing_when_not_due(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config = make_config(schedules={"mode": "weekly", "weekly_day": 1, "time": "10:00"})
    (tmp_path / "social_profiles.json").write_text(json.dumps(config), encoding="utf-8")
    monkeypatch.setattr(tracker, "datetime", FixedDate)
    FixedDate.fixed = datetime(2024, 3, 4)  # Monday, target is Tuesday
    tracker.job_runner()
    assert not Path(tracker.EXCEL_FILE).exists()


# One history row per profile per schedule period

class Scraper:
    """Fake scraper returning queued results and counting calls."""

    def __init__(self, *results) -> None:
        self.results = list(results)
        self.calls = 0

    def __call__(self, url, **kwargs) -> dict:
        self.calls += 1
        return {"username": "a", "followers": self.results.pop(0)}


def run_job(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, mode: str, when: datetime,
            scraper: Scraper, weekly_day: int = 0) -> None:
    config = make_config(schedules={"mode": mode, "weekly_day": weekly_day, "time": "23:00"})
    (tmp_path / "social_profiles.json").write_text(json.dumps(config), encoding="utf-8")
    monkeypatch.setattr(tracker, "SCRAPERS", {"instagram": scraper})
    monkeypatch.setattr(tracker, "datetime", FixedDate)
    FixedDate.fixed = when
    tracker.job_runner()


def test_monthly_startup_and_scheduled_run_record_one_row(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    scraper = Scraper(100, 105)
    run_job(tmp_path, monkeypatch, "monthly", datetime(2024, 2, 29, 9, 0), scraper)  # startup
    run_job(tmp_path, monkeypatch, "monthly", datetime(2024, 2, 29, 23, 0), scraper)  # scheduled
    rows = read_rows(tracker.EXCEL_FILE)
    assert len(rows) == 2 and rows[1][3] == 100
    assert scraper.calls == 1  # the second run did not even scrape


def test_monthly_next_period_records_again(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    scraper = Scraper(100, 110)
    run_job(tmp_path, monkeypatch, "monthly", datetime(2024, 2, 29, 23, 0), scraper)
    run_job(tmp_path, monkeypatch, "monthly", datetime(2024, 3, 31, 23, 0), scraper)
    rows = read_rows(tracker.EXCEL_FILE)
    assert [r[3] for r in rows[1:]] == [100, 110]
    assert rows[2][4] == "+10"


def test_weekly_restart_on_scheduled_weekday_records_one_row(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    scraper = Scraper(1, 2, 3)
    tue = datetime(2024, 3, 5, 9, 0)
    run_job(tmp_path, monkeypatch, "weekly", tue, scraper, weekly_day=1)
    run_job(tmp_path, monkeypatch, "weekly", tue.replace(hour=23), scraper, weekly_day=1)
    assert len(read_rows(tracker.EXCEL_FILE)) == 2
    run_job(tmp_path, monkeypatch, "weekly", datetime(2024, 3, 12, 23), scraper, weekly_day=1)
    assert [r[3] for r in read_rows(tracker.EXCEL_FILE)[1:]] == [1, 2]


def test_daily_restart_records_one_row_per_day(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    scraper = Scraper(1, 2, 3)
    run_job(tmp_path, monkeypatch, "daily", datetime(2024, 3, 5, 9), scraper)
    run_job(tmp_path, monkeypatch, "daily", datetime(2024, 3, 5, 23), scraper)
    run_job(tmp_path, monkeypatch, "daily", datetime(2024, 3, 6, 23), scraper)
    assert [r[3] for r in read_rows(tracker.EXCEL_FILE)[1:]] == [1, 2]


def test_error_row_is_replaced_by_later_success_in_same_period(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    scraper = Scraper(None, 120)
    run_job(tmp_path, monkeypatch, "monthly", datetime(2024, 2, 29, 9), scraper)
    assert read_rows(tracker.EXCEL_FILE)[1][3] == "Error"
    run_job(tmp_path, monkeypatch, "monthly", datetime(2024, 2, 29, 23), scraper)
    rows = read_rows(tracker.EXCEL_FILE)
    assert len(rows) == 2
    assert rows[1][1] == "23:00:00" and rows[1][3] == 120


def test_repeated_error_keeps_single_error_row_and_success_is_not_overwritten(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    scraper = Scraper(None, None, 5)
    for hour in (8, 9, 10):
        run_job(tmp_path, monkeypatch, "monthly", datetime(2024, 2, 29, hour), scraper)
    rows = read_rows(tracker.EXCEL_FILE)
    assert len(rows) == 2 and rows[1][3] == 5


def test_existing_history_from_earlier_period_and_manual_dates_still_append(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Instagram"
    ws.append(["Date", "Time", "Username", "Followers", "Change"])
    ws.append(["2024-01-31", "23:00:00", "a", 90, None])
    ws.append(["not a date", "x", "a", 91, None])
    wb.save(tracker.EXCEL_FILE)
    run_job(tmp_path, monkeypatch, "monthly", datetime(2024, 2, 29, 23), Scraper(100))
    rows = read_rows(tracker.EXCEL_FILE)
    assert len(rows) == 4 and rows[3][3] == 100 and rows[3][4] == "+9"
